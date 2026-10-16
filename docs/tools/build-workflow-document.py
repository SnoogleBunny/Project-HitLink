"""Build the replacement Word specification from the authoritative Markdown.

Run with bundled Python and python-docx. An optional --import-v12 source.docx
imports the retained specification once; normal builds need only the Markdown.
"""
from pathlib import Path
import argparse
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOCS = Path(__file__).resolve().parents[1]
SOURCE = DOCS / 'approved-workflows.md'
OUTPUT = DOCS / 'HitLink_Product_Workflows_Shipping_Specification_v1_3.docx'

UPDATES = '''## Approved product experience

This edition replaces version 1.2 and incorporates the approved interactive mockup. The workflow inventory, domain safeguards and release gates below remain required. This section defines the approved navigation, presentation and connected actions. Where an older roadmap or document conflicts, these explicit decisions take precedence. A rendered screen alone is not evidence of a working workflow.

### Delivery scope and source of truth

Implement the experience in the real owner and member web applications using persisted workspace records, protected server operations and the shared domain rules. Preserve the one-location modular monolith, Next.js, Prisma, PostgreSQL, Stripe Connect and email-only launch architecture. The three account roles remain owner, coach and customer; functional reception, guardian, payer and support permissions are grants within that model, not implied access to every record.

The approved scope covers the mockup's owner, gym member, guardian, adult coaching client, coach and public visitor journeys. Existing P0 and enabled P0C requirements are unchanged. Public paid membership signup is now included in this implementation, superseding its former P1 sequencing. Optional appointment and rank modules remain conditional on the service offered. Deferred items in Appendix A remain deferred.

Mockup samples are fixtures only. Do not ship the single-active-booking restriction, fixed October 2026 date range, hardcoded people or balances, browser-memory persistence, simulated payments, simulated notices, scenario controls or reset tools as production behavior. Multiple valid bookings per person must work subject to the accepted capacity, conflict, entitlement and cutoff rules. Use the workspace clock and timezone. Real actions must survive reloads and independent sessions. The compact preview frame is not a fixed production page height; use the actual viewport with stable navigation and scrollable content.

The public product name is HitLink. Technical package identifiers may remain unchanged where renaming would add unrelated risk. Approval of this specification authorizes implementation; it does not certify production activation or waive verification of money, privacy, migration and recovery.

### Navigation and record ownership

Owner primary destinations are Today, Classes, Members, Billing, Plans, Forms, Website and Settings. Detail routes keep their parent selected and provide a clear return path. Retain useful alerts and direct links to exact records rather than embedding every detail workflow in an overview page. Preserve selected class, date and member when navigating or refreshing a valid deep link; invalid or missing drafts return safely to their parent.

Classes owns dated rosters, class creation, one-date edits, future-series edits and coach assignments. Members owns member details, membership changes, family permissions and child rank history. Settings owns program-wide rank rules, Import and switch systems, Records and account closure, and Support and audit. Import review and cutover are stages of the same workflow, retaining separate reconciliation, payer authorization, owner acknowledgment and operator activation gates.

The member portal primary destinations are Classes, Membership, Payments and Account. Classes is the gym member landing page and combines booking discovery with management. There is no additional gym My day, Home or My bookings dashboard. Legacy routes redirect to their current destinations without losing an authorized booking context. Account owns household relationships; Membership links to Agreements and forms, and remains selected on signing routes. Payments is signed-in payer scoped, independent of the person selected for classes.

Coaches start at Clients. A selected client has Plan, Notes and, when the program includes them, Appointments sections. Enrollment and accepted records belong to that client's plan. Pause, replacement and closure are focused program actions. Notes-only programs omit appointment navigation and resolve old appointment links to the program. Coaching clients retain My day for daily notes, My program and, when offered, My appointment. Daily notes and plan approvals are distinct tasks.

### Owner Today and operational cards

Today starts with operational information and a member-portal link, without a duplicate introductory card. Place trial follow-up prominently above lower-priority content. Each metric names its population and destination. Active members uses real active records and opens the complete member directory, clearing any previous family or status filter. Its destination label says View all members. Class bookings names the class, shows booked places against capacity and opens that dated roster. The displayed class and count must refer to the same occurrence.

Payments needing attention opens Billing with its attention filter. Count failed or unresolved payments, pending refund recovery and unresolved dispute cases consistently with the queue and alerts. Assigning a reviewer does not settle a dispute. Only a verified outcome closes a financial obligation. At zero, the card remains readable without an action, chevron, keyboard stop or hover affordance. Peer metric cards share dimensions and aligned content.

### Classes and recurring schedule changes

Owners add classes from Classes and edit from the exact dated roster. A focused form captures title, program or class type, date, local start and end, room, capacity, optional default coach and applicable booking/cancellation settings. Support one date, daily recurrence and selected weekdays. Preview resolved occurrence dates in the workspace timezone, including daylight-saving issues, before saving.

Edits explicitly target this class only or this and future classes. Show the affected dates, bookings and waitlist entries before confirmation. Protect historical and completed attendance, independent date exceptions and stable occurrence IDs. Reject coach or room overlaps, capacity below confirmed bookings, unresolved payment-pending reservations, and removal of booked or waitlisted dates unless the supported cancellation/resolution workflow is used. Revalidate on commit so a stale preview cannot overwrite concurrent bookings or edits.

An occurrence's room appears directly beneath its class title in owner Today, owner Classes and member class listings. Coach identity and Assign coach or Change coach remain together in a separate area. Booking reviews explicitly name class, local date/time, participant, room, current coach, funding source and consequences.

The representative seed timetable repeats daily: HIIT Bootcamp 7:15 am to 8:30 am; HIIT/Hyrox 4:30 pm to 5:30 pm; Beginner Padwork 5:30 pm to 6:30 pm; Beginner Muay Thai 6:30 pm to 7:30 pm; Advanced Muay Thai 6:30 pm to 7:30 pm. The simultaneous classes remain distinct records with separately valid resources. These are demo seeds, not immutable business settings; owners can create their actual timetable.

### Coach assignment by date

A recurring class may have no default coach. Each occurrence independently inherits the default, explicitly assigns a coach or explicitly remains unassigned. A one-date assignment does not create a complete schedule exception: later time, room or capacity changes still apply unless those fields have their own dated exception. Changing the default preserves explicit assigned and unassigned dates. Historical teacher records and recorded attendance are immutable except through an audited correction.

Today, Classes and the selected roster provide a quick Assign coach or Change coach action for the exact date. Owners see Needs coach; members see Coach to be confirmed and may still book when otherwise eligible. Persist the assignment mode as well as the effective coach so changing inheritance without changing the displayed coach is not lost.

Check the effective coach against other classes, enabled coaching appointments, leave and applicable buffers. Unassigned values do not conflict with one another. These checks detect scheduling conflicts; they do not verify qualifications or actual willingness to teach. Staff confirm availability. A changed published coach durably queues notices for affected booked and waitlisted people; changing assignment mode with no published-name change requires no notice. Preserve seats, attendance provenance, money, funding and allowance.

### Unified member Classes

Show the selected day and program filter alongside that member's exact Booked, Pending, Waitlisted and Attended states. A Manage action opens the relevant booking's receipt, move or cancellation controls. Keep compact attendance history available. If the current day or filter hides an active booking, show a named link that selects its date and clears the incompatible filter. Empty states offer direct date selection and booking rather than a dead end.

Booking confirmation, uncertain-result recovery, waitlist promotion and moves return to the matching class/date/person. A move rechecks the target and transfers the original hold atomically before releasing the original seat; failure retains the original. Past and future dates with the same weekday remain distinct. Allow multiple active reservations where policy permits; the mockup's one-reservation simplification is not a business rule.

Use one compact text summary for each finite allowance period: remaining, reserved, used and exact reset date. Show weekly and billing-month limits separately; both constrain the same reservation and must never be added. Do not use a bar, dots or decorative consumption meter. Unlimited plans omit the allowance panel and state Unlimited classes in membership details. For a single finite cap, show only that period. Counting policy remains owner configurable; booking and attendance never count one class twice. Put detailed policy in a disclosure.

Show waiver and unresolved-renewal reminders only when relevant to the selected beneficiary. Verified settlement clears only its resolved warning. Membership contains accepted plan, billing dates, freeze/cancellation status and forms. Omit a punch-card panel when no such entitlement exists.

### Family identity and consent

Use one compact person picker in the member header on Classes, Membership and Forms. Offer selection only when more than one verified authorized person is available. Desktop uses an anchored picker; phone uses an accessible sheet. Preserve the selected person across navigation, reset incompatible filters and show a skeleton while their records load. Avoid repeating a dedicated profile card or name banner on every page.

Changing the beneficiary never changes the signed-in account or payer scope. Keep household/account pages explicitly scoped. Name the participant in consequential booking, payment and signing confirmations. Unverified or revoked links must not appear or succeed through crafted requests. Prevent background switching while a booking/signing dialog is open; an unfinished waiver requires explicit discard before switching. Each person's reservations, consent, quotas and records remain separate.

### Membership policy and billing

Preserve the full billing-cycle allowance, notice, freeze and payment rules in Section 1. Owners configure reservation versus attendance presentation, penalties, weekly/monthly caps, finite freezes and notice policy. The recommended two-calendar-month notice policy is a configurable default, not a mandatory commercial term. Preview final planned renewal, access end, affected bookings, refund/credit consequences and recorded owner exceptions before approval. A settings change cannot rewrite an accepted contract.

Billing offers actionable records first with an explicit history view and an attention filter. Every row identifies the payer, beneficiary or service where permitted, obligation, state, amount, currency, deadline and available action. Keep the currency code and amount in one nonbreaking text run with a narrow space, tabular figures and two decimal places, including zero. CAD is the Canadian demo format; use the actual workspace or accepted transaction currency in live records.

Update card, retry payment, refund, credit, dispute review and service closure are separate operations. Never treat a card update, redirect, button click or local success toast as verified settlement. Unknown outcomes stay pending with a recovery path. Refund review reserves refundable money and affected entitlement; concurrent requests cannot exceed the remaining amount. Dispute reviewer assignment stays open until a verified outcome. Real provider callbacks, persisted operations, reconciliation and durable notices replace all simulated mockup actions.

### Adult coaching programs and notes

Deliver structured adult one-to-one programs with accepted plan versions, milestones, timezone, inclusive delivery dates, note calendar, response terms, coach assignment and upfront payment. Zero-appointment programs are fully usable. Appointments appear only when sold and must use the shared conflict and unit rules. Notes never consume sessions or money. Client journals, interactive check-ins and automated coaching remain deferred.

The coach's Notes area has distinct author-private working notes and daily client notes, with unmistakable audience labels. Saving a draft never publishes it. Publishing previews named client, enrollment, program date and content, then commits a version and queues a minimal secure-link notice. Clients see only published versions and history in My day, with due, scheduled, missed-delivery, amended and withdrawn states. Actual publication time remains distinct from the program date. Stale saves and duplicate requests cannot overwrite newer versions or double-publish.

Only the authorized current author may access private working notes or unpublished drafts. Ownership, reception, payer, former guardian and unrelated staff roles confer no content access. Private working notes are excluded from every product export, including owner, client, handoff, offboarding, incident and reviewed-request export jobs. A formal request requiring exceptional content review uses a separately authorized, time-limited and audited read process; it does not unlock a bulk export path. Published client records may be exported only within the named client's authorized scope. Notifications, logs, analytics and ordinary search omit restricted text.

Finite pause, awaiting replacement, accepted handoff, completion, extension and early closure retain the Section 1 consequences. Revoke the outgoing coach immediately. A new coach receives only client-approved published history and shared plan; private notes never transfer. Service closure, refund settlement and record retrieval have independent status. No action changes a gym membership or recreates old reminders. Mistaken publication uses the restricted incident and withdrawal workflow.

### Public website signup and legal pages

Owner Website contains Pricing and Search and sharing. Persist pricing drafts across navigation. Publish explicitly after validation; ordinary visitors receive the last published version. One visible plan may be highlighted or none; hiding it clears the highlight. Draft previews are labelled and cannot start enrollment. Each published offer must resolve to confirmed operational terms and an exact offer version before accepting payment; marketing copy never silently alters an accepted contract.

Visitors see Memberships and Try a class in a public header without application sidebars, avatars or member tabs. Place the trial entry prominently. Joining continues from the selected plan. Keep trial and paid signup separate and preserve the selected offer and entered participant through account verification, required agreements, hosted payment and inline confirmation. Adult and guardian paths maintain distinct identities and authority. A single trial option is a static summary. Marketing consent is optional and separate.

Provide Member portal, HitLink terms and Privacy links in the footer at phone and desktop sizes. Legal navigation preserves in-progress form details. Terms and Privacy describe the HitLink software service, with clearly marked business-name, contact and operational placeholders until supplied. Placeholder legal text remains visibly draft and cannot be represented as final legal readiness. The custom missing-page route returns a real HTTP 404 and offers useful recovery links.

Public routes have unique titles, descriptions, canonical metadata and share images. Private routes omit personal data from metadata and remain protected by authorization; robots directives are not access control. Draft and preview pages use noindex. Website metadata settings validate allowed public fields and never expose coaching or member records. Gym offers are distinct from HitLink software subscription pricing.

### Visual system and responsive behavior

Use the approved neutral light/dark surfaces and forest-green action accent. Light canvas/navigation/cards are #F3F3F4, #FBFBFB and #FFFFFF; dark equivalents are #2C2D32, #323339 and #36373E. Primary accent is #15392B with a sufficiently contrasting related edge/text treatment in dark mode. Use Barlow Semi Condensed 600/700 for major headings and prominent values and Source Sans 3 for navigation, controls, tables and notes. Do not use Inter. Load paired fonts together with a bounded fallback; font loading must not block the app indefinitely.

Use shared components for cards, buttons, fields, badges, banners, money and skeletons. Short status badges are consistently 24px high. Peer cards align outer dimensions and their decision/action areas; stacked phone cards keep natural content height. Use real useful status and actions rather than filler, repeated content or invented trends. Keep terminal actions at the bottom with at least 16px separation; actions inside records remain with those records.

Apply a 4px spacing scale: 8px for grouped labels/actions, 12px for related content, 16px for card content/forms, 20px card padding and 24px gutters/sections on desktop. Phones use 16px insets and 20px sections. Layout gaps own spacing instead of accumulating child margins. Banners put actions in a separate row 12px below the message. Room text sits under class titles; currency and figures stay together. Use plain task-specific language without slogans or internal implementation explanations.

Buttons have obvious hover fills, contrast and pressed states without moving or resizing. Text links use underlines and color feedback, without filled-button hover effects. Align card/list text links with surrounding text using zero horizontal padding and a compact 28px desktop target; phone/coarse targets remain at least 44px. Icon-only actions retain button feedback. Focus outlines and disabled states remain visible. Dropdowns use consistent 12px text insets, separate arrow space, wrap long selected values and preserve native keyboard behavior and fallbacks.

Every loading destination, person switch and record fetch uses a structure-matched skeleton, with a truthful loading status and stable layout. Do not hide errors behind perpetual skeletons or fake delays. Use immediate committed-state updates or pending indicators for mutations; show success only after the supported server result. A theme toggle follows system preference initially, persists a deliberate choice and changes no route, scroll, draft or loading state.

At 640px and below, use a compact 56px header, persistent common navigation and a scrollable full menu. Owner phone tabs are Today, Classes and Members plus More; member tabs are Classes, Membership and More. Coach Clients uses header/local navigation without a redundant single-tab bar. Reflow tables into labelled records without dropping fields. Use 44px targets and 16px editable text. Seven date choices fit one row from 390px and wrap on narrower phones. Dialog content scrolls while heading and actions remain available. Menus and dialogs isolate and restore focus. Feedback occupies its own space above bottom navigation and cannot cover records or controls.

Use short purposeful route, dialog and action-result motion consistent with the approved mockup. Animate destinations as one surface, not staggered cards or intermediate balances. Controls become usable as soon as loading finishes. Cancel stale callbacks on rapid navigation, overlay replacement and breakpoint changes; pause motion in hidden documents. Reduced motion uses immediate changes with the same text, focus and announcements.

### Acceptance evidence for the approved implementation

Verify connected workflows against persisted data in separate owner, member, guardian, coach and client sessions. Cover owner metric destinations, exact dated rosters, add/edit class scopes, all three coach-assignment modes, unified booking/move/cancel/waitlist/recovery, independent family quotas and signing rights, configurable membership requests, real payment recovery, website draft/publish/signup, notes-only coaching and accepted appointment programs.

Exercise negative paths: stale schedule edits, racing last seat/allowance/unit, duplicate provider events, pending refunds, revoked family links/coaches, wrong client or workspace, publication race, private-note export attempts and unsupported service activation. Demonstrate rollback or reconciliation without deleting accepted history. Existing workflow gates in Sections 3 to 6 remain the release standard; this checklist does not replace them.

Check desktop, tablet and 320px/390px phone layouts in both themes, keyboard-only use, focus restoration, reduced motion, loading/failure states, long labels and real amounts. Confirm no horizontal overflow, hidden controls, mismatched card/badge sizes or currency gaps. Build, typecheck, database migrations and meaningful automated tests must pass. Preserve source-controlled acceptance evidence and distinguish implemented behavior from remaining operational configuration or release gates.

'''

def revised(text):
    text = text.replace('Version 1.2  |  October 9, 2026', 'Version 1.3  |  October 15, 2026')
    text = text.replace('Owner defines name, program, duration, capacity, room and default coach.', 'Owner defines name, program, duration, capacity, room and an optional default coach; dated assignments remain independent of schedule exceptions.')
    text = text.replace('Substitution is visible to staff and members and preserves the original series configuration.', 'Each occurrence inherits the optional default coach, explicitly assigns a substitute, or explicitly remains unassigned. Changes preserve other schedule inheritance, reservations and attendance history and notify affected customers when the published coach changes.')
    text = text.replace('Program, room and coach exist; recurring series is published.', 'Program and required room exist; recurring series is published. Default coach is optional, with effective dated assignments checked for conflicts.')
    text = text.replace('View personal dashboard', 'Use the role appropriate portal landing page')
    text = text.replace('A separate responsive member/guardian/client portal shows only authorized next actions, selected beneficiary, bookings and allowance. Hybrid users see distinct gym memberships and coaching programs.', 'A separate responsive portal opens gym members and guardians on Classes with unified booking management and a compact authorized-person header picker. Coaching clients use My day for published daily notes. Hybrid users retain separate memberships and coaching programs.')
    text = text.replace('Private working notes use reviewed access; other services continue.', 'Private working notes remain excluded from every product export; exceptional read access requires a separate authorized review. Other services continue.')
    text = text.replace('Not in ordinary portal export', 'Excluded from every product export')
    text = text.replace('routine exports/logs omit private text.', 'every product export and routine logs omit private text.')
    text = text.replace('General exports exclude restricted coaching content.', 'General exports exclude restricted coaching content. Private working notes are excluded from every product export, including reviewed requests; exceptional read access remains a separate authorized process.')
    text = text.replace('Suggested concurrency rule for the agent team', 'Delivery coordination')
    return text

def import_v12(path):
    d = Document(path)
    chunks = ['# HitLink Product Workflows and Shipping Specification', '', 'Martial arts and adult life coaching', '', 'Version 1.3  |  October 15, 2026', '', 'Prepared for the founder and product delivery team.', '', 'This specification defines the approved product experience, domain workflows and release acceptance for HitLink. It preserves the version 1.2 workflow inventory and incorporates the approved mockup decisions. It is a delivery requirement, not a claim that every requirement is already implemented.', '', UPDATES]
    started = False
    for el in d.element.body:
        if el.tag == qn('w:p'):
            p = Paragraph(el, d)
            text = revised(p.text.strip())
            if text == 'Scope of this edition': started = True
            if not started or not text: continue
            if p.style.name.startswith('Heading'):
                level = 3 if text.startswith(('CF ', 'Epic ')) else min(int(p.style.name[-1]) + 1, 4)
                chunks += ['#' * level + ' ' + text, '']
            else:
                chunks += [text.replace('\n', '  \n'), '']
        elif el.tag == qn('w:tbl') and started:
            t = Table(el,d)
            rows = [[revised(c.text.strip()) for c in row.cells] for row in t.rows]
            if len(rows)==1 and len(rows[0])==1:
                parts=rows[0][0].split('\n',1)
                if len(parts)>1 and parts[0] in ['Scheduling architecture decision','Delivery coordination','Release test']:
                    chunks += [parts[1], '']
                else:
                    chunks += [('**'+parts[0]+'** '+parts[1]) if len(parts)>1 else parts[0], '']
            elif rows[0][0] in ['Outcome', 'Workflow field']:
                if rows[0][0]=='Workflow field': rows=rows[1:]
                for label, text in rows:
                    chunks += ['**'+label+'** '+text, '']
            else:
                # Explicit approval promotes public paid enrollment in this edition.
                for row in rows:
                    if row[0]=='MP-06': row[3]='P0'
                chunks += ['| '+' | '.join(c.replace('|','\\|').replace('\n','<br>') for c in rows[0])+' |', '| '+' | '.join('---' for _ in rows[0])+' |']
                chunks += ['| '+' | '.join(c.replace('|','\\|').replace('\n','<br>') for c in row)+' |' for row in rows[1:]]
                chunks += ['']
    SOURCE.write_text('\n'.join(chunks),encoding='utf-8')

def add_runs(p,text):
    for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        run=p.add_run(part.replace('<br>','\n').replace('\\|','|'))
        run.bold=i%2==1

def set_cell(cell, text, header, row_index):
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=cell.paragraphs[0]
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=1.05
    add_runs(p,text)
    for r in p.runs:
        r.font.size=Pt(8.5)
        r.font.color.rgb=RGBColor.from_string('FFFFFF' if header else '000000')
        if header:r.bold=True
    tcPr=cell._tc.get_or_add_tcPr()
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'333333' if header else ('F3F3F3' if row_index%2==0 else 'FFFFFF'));tcPr.append(shade)
    margins=OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
    tcPr.append(margins)

def add_table(d,rows):
    n=len(rows[0]);t=d.add_table(rows=0,cols=n)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths={5:[.54,1.23,.87,.43,4.07],3:[1.35,2.0,3.79],2:[1.55,5.59]}.get(n,[7.14/n]*n)
    if rows[0][0]=='Record':widths=[1.12,1.43,1.65,1.53,1.41]
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci,text in enumerate(row):
            cells[ci].width=Inches(widths[ci]);set_cell(cells[ci],text,ri==0,ri)
            if (n==5 and ci in [0,3] and rows[0][0]=='ID'):cells[ci].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        trPr=t.rows[-1]._tr.get_or_add_trPr()
        noSplit=OxmlElement('w:cantSplit');trPr.append(noSplit)
        if ri==0:
            header=OxmlElement('w:tblHeader');trPr.append(header)
            for c in cells:c.paragraphs[0].paragraph_format.keep_with_next=True
    d.add_paragraph().paragraph_format.space_after=Pt(2)

def build():
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.58);sec.bottom_margin=Inches(.58);sec.left_margin=Inches(.68);sec.right_margin=Inches(.68)
    for name,size in [('Normal',9.5),('Title',25),('Subtitle',12),('Heading 1',18),('Heading 2',13),('Heading 3',10.5)]:
        s=d.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_after=Pt(6)
        s.paragraph_format.line_spacing=1.08
        if name.startswith('Heading'):
            s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(12)
    d.styles['Normal'].paragraph_format.widow_control=True
    for name in ['Heading 1','Heading 2','Heading 3','Title']:
        ppr=d.styles[name].element.get_or_add_pPr()
        for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before=Pt(0)
    r=footer.add_run('HitLink 1.3  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    lines=SOURCE.read_text(encoding='utf-8').splitlines();i=0;section_count=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            rows=[]
            while True:
                if not re.match(r'^\|[\s:|\-]+\|$',line):rows.append([c.strip() for c in re.split(r'(?<!\\)\|',line)[1:-1]])
                if i>=len(lines) or not lines[i].startswith('|'):break
                line=lines[i].strip();i+=1
            add_table(d,rows);continue
        if line.startswith('#'):
            marker,title=line.split(' ',1);level=len(marker)
            # Headings use plain words and numbers, while body preserves source punctuation.
            title=re.sub(r'[^\w\s]',' ',title);title=re.sub(r'\s+',' ',title).strip()
            p=d.add_paragraph(style='Title' if level==1 else 'Heading '+str(min(level-1,3)))
            add_runs(p,title)
            if title.startswith(('CF ', 'Epic ')):
                for run in p.runs:run.font.size=Pt(10.5)
            if level==2:
                p.paragraph_format.page_break_before=section_count>0 and not re.match(r'^[567A] ',title)
                section_count+=1
            continue
        p=d.add_paragraph();add_runs(p,line)
        if line.startswith('☐'):p.paragraph_format.left_indent=Inches(.15);p.paragraph_format.first_line_indent=Inches(-.15)
    d.core_properties.title='HitLink Product Workflows and Shipping Specification'
    d.core_properties.subject='Approved martial arts and adult life coaching workflows'
    d.core_properties.author='HitLink'
    d.core_properties.version='1.3'
    d.save(OUTPUT)
    print(f'Built {OUTPUT} from {SOURCE}')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--import-v12',type=Path);parser.add_argument('--markdown-only',action='store_true');args=parser.parse_args()
    if args.import_v12:import_v12(args.import_v12)
    if not args.markdown_only:build()
